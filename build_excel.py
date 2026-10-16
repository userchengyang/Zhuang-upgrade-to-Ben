# -*- coding: utf-8 -*-
"""
把 scraper.py 抓到并缓存的数据整理成 Excel。

用法：
  python build_excel.py                       # 读取 cache/ 生成 河南专升本数据.xlsx
  python build_excel.py --cache cache --out 河南专升本数据.xlsx

生成的工作表：
  录取分数线  —— 每个「专业 + 院校」的历年录取分数
  招生计划    —— 每个「专业 + 院校」的招生计划、学制、学费、住宿费等
  省控线      —— 每个专业历年的省控线及对应招生人数
"""
import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import scraper

DEFAULT_CACHE = "cache"
DEFAULT_OUT = "河南专升本数据.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF")
CENTER = Alignment(horizontal="center", vertical="center")


def load_all(cache_dir):
    items = []
    for fn in sorted(os.listdir(cache_dir)):
        if not fn.endswith(".json"):
            continue
        with open(os.path.join(cache_dir, fn), "r", encoding="utf-8") as f:
            try:
                items.append(json.load(f))
            except Exception:
                pass
    return items


def collect_years(items, kind):
    years = set()
    for it in items:
        data = it.get(kind)
        if not data:
            continue
        if kind == "schoolScore":
            for school in data.get("plan") or []:
                for entry in school.get("score") or []:
                    if ":" in entry:
                        years.add(entry.split(":", 1)[0])
        elif kind == "provinceScore":
            for row in data.get("data") or []:
                y = str(row.get("year", "")).strip()
                if y:
                    years.add(y)
    return sorted(years, key=lambda x: int(x) if x.isdigit() else 0)


def write_sheet(ws, header, rows):
    ws.append(header)
    for row in rows:
        ws.append(row)
    for c in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(header)), ws.max_row)


def autosize(ws, max_width=60):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            v = "" if cell.value is None else str(cell.value)
            w = sum(2 if ord(ch) > 127 else 1 for ch in v)
            widths[cell.column] = max(widths.get(cell.column, 0), w)
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(w + 2, 8), max_width)


def build(cache_dir, out_path):
    items = load_all(cache_dir)
    items.sort(key=lambda it: it.get("major", ""))
    print("载入专业数:", len(items))

    wb = Workbook()

    # 1) 录取分数线
    years = collect_years(items, "schoolScore")
    ws1 = wb.active
    ws1.title = "录取分数线"
    header1 = ["专业", "院校", "类型", "院校代码"] + years
    rows1 = []
    for it in items:
        major = it.get("major", "")
        data = it.get("schoolScore") or {}
        schools = data.get("plan") or []
        schools = sorted(schools, key=lambda s: s.get("name", ""))
        for s in schools:
            smap = {}
            for entry in s.get("score") or []:
                if ":" in entry:
                    y, v = entry.split(":", 1)
                    smap[y.strip()] = v.strip().replace("\t", "")
            row = [major, s.get("name", ""), s.get("type", ""), s.get("code", "")]
            for y in years:
                if y in smap:
                    row.append(smap[y] if smap[y] != "" else "/")
                else:
                    row.append("×")
            rows1.append(row)
    write_sheet(ws1, header1, rows1)
    autosize(ws1)

    # 2) 招生计划
    ws2 = wb.create_sheet("招生计划")
    header2 = ["专业", "院校", "类型", "学制", "学费(元)", "住宿费(元)", "招生计划(人)", "地址/备注"]
    rows2 = []
    for it in items:
        major = it.get("major", "")
        data = it.get("plan") or {}
        plans = data.get("plan") or []
        plans = sorted(plans, key=lambda p: p.get("school", ""))
        for p in plans:
            addr = p.get("remarks") or p.get("address") or ""
            rows2.append([
                major,
                p.get("school", ""),
                p.get("type", ""),
                p.get("length", ""),
                p.get("study_fee", ""),
                p.get("room_fee", ""),
                p.get("plan", ""),
                addr,
            ])
    write_sheet(ws2, header2, rows2)
    autosize(ws2)

    # 3) 省控线
    pyears = collect_years(items, "provinceScore")
    ws3 = wb.create_sheet("省控线")
    header3 = ["专业"] + ["%s分数" % y for y in pyears] + ["%s人数" % y for y in pyears]
    rows3 = []
    for it in items:
        major = it.get("major", "")
        data = it.get("provinceScore") or {}
        smap, nmap = {}, {}
        for row in data.get("data") or []:
            y = str(row.get("year", "")).strip()
            smap[y] = row.get("score", "")
            nmap[y] = row.get("nums", "")
        r = [major]
        r += [smap.get(y, "×") for y in pyears]
        r += [nmap.get(y, "×") for y in pyears]
        rows3.append(r)
    write_sheet(ws3, header3, rows3)
    autosize(ws3)

    wb.save(out_path)
    print("已生成:", out_path)
    print("  录取分数线: {} 行".format(len(rows1)))
    print("  招生计划  : {} 行".format(len(rows2)))
    print("  省控线    : {} 行".format(len(rows3)))


def main():
    ap = argparse.ArgumentParser(description="整理河南专升本数据为 Excel")
    ap.add_argument("--cache", default=DEFAULT_CACHE, help="缓存目录")
    ap.add_argument("--out", default=DEFAULT_OUT, help="输出 Excel 文件名")
    args = ap.parse_args()
    if not os.path.isdir(args.cache):
        print("找不到缓存目录:", args.cache, "，请先运行 python scraper.py --all")
        return
    build(args.cache, args.out)


if __name__ == "__main__":
    main()
